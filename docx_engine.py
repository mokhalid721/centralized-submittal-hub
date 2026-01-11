import re
from io import BytesIO
from pathlib import Path
from typing import Dict

from docx import Document

PLACEHOLDER_RE = re.compile(r"«([^»]+)»")  # captures inside «...»
WHOLE_PLACEHOLDER_PARA_RE = re.compile(r"^\s*(?:«[^»]+»\s*)+$")
BULLET_ONLY_RE = re.compile(r"^[•\-\u2013\u2014]\s*$")  # • - – —


def iter_paragraphs(container):
    for p in getattr(container, "paragraphs", []):
        yield p
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def iter_all_paragraphs(doc: Document):
    yield from iter_paragraphs(doc)
    for sec in doc.sections:
        yield from iter_paragraphs(sec.header)
        yield from iter_paragraphs(sec.footer)
        yield from iter_paragraphs(sec.first_page_header)
        yield from iter_paragraphs(sec.first_page_footer)
        yield from iter_paragraphs(sec.even_page_header)
        yield from iter_paragraphs(sec.even_page_footer)


def extract_placeholders_from_docx(path: str | Path):
    doc = Document(str(path))
    found = set()
    for p in iter_all_paragraphs(doc):
        for name in PLACEHOLDER_RE.findall(p.text):
            found.add(name.strip())
    return sorted(found)


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)


def replace_token_across_runs(paragraph, token: str, replacement: str):
    """
    Replace token across Word runs while preserving surrounding formatting.
    Word often splits text into separate runs when formatting changes.
    """
    if not paragraph.runs:
        return

    while True:
        full = "".join(r.text for r in paragraph.runs)
        idx = full.find(token)
        if idx == -1:
            break

        start = idx
        end = idx + len(token)

        pos = 0
        start_run = end_run = None
        start_off = end_off = 0

        for i, r in enumerate(paragraph.runs):
            next_pos = pos + len(r.text)

            if start_run is None and start < next_pos:
                start_run = i
                start_off = start - pos

            if start_run is not None and end <= next_pos:
                end_run = i
                end_off = end - pos
                break

            pos = next_pos

        if start_run is None or end_run is None:
            break

        if start_run == end_run:
            r = paragraph.runs[start_run]
            r.text = r.text[:start_off] + replacement + r.text[end_off:]
        else:
            first = paragraph.runs[start_run]
            last = paragraph.runs[end_run]

            first.text = first.text[:start_off] + replacement
            for j in range(start_run + 1, end_run):
                paragraph.runs[j].text = ""
            last.text = last.text[end_off:]


def fill_docx_to_bytes(template_path: str | Path, values: Dict[str, str]) -> BytesIO:
    """
    values keys must be placeholder keys WITHOUT angle markers.
    """
    doc = Document(str(template_path))
    token_map = {f"«{k}»": (values.get(k, "") or "") for k in values.keys()}

    # keep original text to decide if a paragraph was only placeholders
    original_texts = {}
    for p in iter_all_paragraphs(doc):
        original_texts[id(p)] = p.text
        for token, replacement in token_map.items():
            if token in p.text:
                replace_token_across_runs(p, token, replacement)

    # Cleanup pass
    for p in list(iter_all_paragraphs(doc)):
        original = (original_texts.get(id(p)) or "")
        now = (p.text or "").strip()

        # remove bullet-only lines and empty list-style lines
        style_name = ""
        try:
            style_name = (p.style.name or "").lower()
        except Exception:
            style_name = ""

        is_listish = ("list" in style_name) or ("bullet" in style_name)

        if BULLET_ONLY_RE.match((p.text or "").strip()):
            delete_paragraph(p)
            continue

        if now == "" and is_listish:
            delete_paragraph(p)
            continue

        # remove paragraphs that were ONLY placeholders (and now empty)
        if now == "" and WHOLE_PLACEHOLDER_PARA_RE.match(original.strip() or ""):
            delete_paragraph(p)
            continue

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out
